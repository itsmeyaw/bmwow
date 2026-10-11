"""Lossless feedback XLSX-to-CSV export; no enrichment or provider calls.

openpyxl is already the workbook reader; SheetJS is its TypeScript counterpart.
Source locators travel with each classification row, including unlabelled rows.
"""

import csv
import hashlib
import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from bmwow.ontology.models import SpreadsheetLocator

ORIGIN_COLUMN = "__bmwow_origin"
REQUIRED_COLUMNS = {"ID", "Customer Feedback", "Derivate (e-code)", "Country"}


class FeedbackOrigin(BaseModel):
    source_file: str
    source_hash: str = Field(pattern=r"^sha256:[a-f0-9]{64}$")
    locator: SpreadsheetLocator


class FeedbackExport(BaseModel):
    source_hash: str
    sheets: int
    rows: int
    distinct_feedback: int


def export_feedback_workbook(source: Path, output: Path) -> FeedbackExport:
    """Export all classification rows atomically, preserving original strings and locations."""
    digest = "sha256:" + hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    formulas = load_workbook(source, read_only=True, data_only=False)
    rows: list[dict[str, str]] = []
    columns: list[str] | None = None
    selected_sheets = 0
    try:
        for sheet in workbook:
            header_row = None
            headers = None
            for index, cells in enumerate(sheet.iter_rows(max_row=20), 1):
                values = [str(c.value).strip() if c.value is not None else "" for c in cells]
                if REQUIRED_COLUMNS <= set(values):
                    header_row, headers = index, values
                    break
            if headers is None:
                continue
            if any(not h for h in headers) or len(headers) != len(set(headers)):
                raise ValueError("Feedback header contains empty or duplicate columns")
            if ORIGIN_COLUMN in headers:
                raise ValueError("Feedback source uses a reserved provenance column")
            if columns is not None and columns != headers:
                raise ValueError("Feedback worksheet headers differ; export separately")
            columns = headers
            selected_sheets += 1
            raw_sheet = formulas[sheet.title]
            raw_rows = raw_sheet.iter_rows(min_row=header_row + 1, max_col=len(headers))
            value_rows = sheet.iter_rows(min_row=header_row + 1, max_col=len(headers))
            for index, (cells, original) in enumerate(
                zip(value_rows, raw_rows, strict=True), header_row + 1
            ):
                if all(cell.value is None for cell in cells):
                    if any(cell.data_type == "f" for cell in original):
                        raise ValueError("Feedback formula has no cached result")
                    continue
                if any(
                    c.value is None and raw.data_type == "f"
                    for c, raw in zip(cells, original, strict=True)
                ):
                    raise ValueError("Feedback formula has no cached result")
                row = {
                    header: "" if cell.value is None else str(cell.value)
                    for header, cell in zip(headers, cells, strict=True)
                }
                origin = FeedbackOrigin(
                    source_file=str(source),
                    source_hash=digest,
                    locator=SpreadsheetLocator(
                        kind="spreadsheet",
                        sheet=sheet.title,
                        range=f"A{index}:{get_column_letter(len(headers))}{index}",
                    ),
                )
                row[ORIGIN_COLUMN] = origin.model_dump_json()
                rows.append(row)
    finally:
        workbook.close()
        formulas.close()
    if columns is None:
        raise ValueError("No feedback worksheet with the required headers was found")
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=output.parent, prefix=".feedback-", suffix=".csv")
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=[*columns, ORIGIN_COLUMN])
            writer.writeheader()
            writer.writerows(rows)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, output)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return FeedbackExport(
        source_hash=digest,
        sheets=selected_sheets,
        rows=len(rows),
        distinct_feedback=len({r["ID"] for r in rows if r["ID"].strip()}),
    )


def export_feedback_directory(input_dir: Path, output_dir: Path) -> list[FeedbackExport]:
    sources = sorted(input_dir.glob("*_feedback_hackathon.xlsx"))
    if not sources:
        raise ValueError("No feedback workbooks were found")
    return [export_feedback_workbook(p, output_dir / f"{p.stem}.csv") for p in sources]
