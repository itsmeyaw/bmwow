"""Source-specific XLSX reads into v0.2 Evidence, with original cells and unknowns preserved."""

from __future__ import annotations

import hashlib
import math
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from bmwow.ontology import GraphBundle, validate_bundle
from bmwow.ontology.models import Context, Edge, Evidence, Measurement, Signal, SourceDocument

from .config import KnowledgeConfig
from .manifest import CurationManifest
from .repository import InMemoryRepository
from .service import KnowledgeService


class ImportReport(BaseModel):
    sales_measurements: int = 0
    study_means: int = 0
    study_shares: int = 0
    normalized_totals: int = 0
    missing_cells: int = 0
    reviewed_study_findings: int = 0
    unmapped_questions: int = 0
    # Exact original labels remain private; print only the counts in CLI output.
    questions: list[str] = Field(default_factory=list)


class WorkbookImporter:
    def __init__(
        self,
        manifest: CurationManifest,
        *,
        as_of: date,
        origin: Literal["provided_bmw", "synthetic"] = "provided_bmw",
    ) -> None:
        self.manifest = manifest.model_copy(deep=True)
        self.config = self.manifest.knowledge.model_copy(deep=True)
        self.bundle = manifest.seed(as_of, "fixture" if origin == "synthetic" else "production")
        self.origin = origin
        self.created_at = datetime.combine(as_of, datetime.min.time(), tzinfo=UTC)
        self.report = ImportReport()
        self._questions: set[str] = set()

    def _model(self, code: str) -> str:
        ids = [
            node.id
            for node in self.manifest.context_nodes
            if node.type == "VehicleModel"
            and node.properties.generation_code == code
            and node.properties.variant is None
        ]
        if len(ids) != 1:
            raise ValueError(f"Configure one canonical model for source code {code}")
        return ids[0]

    def _context(self, models: list[str], dataset: str, market: str) -> Context:
        market_id = self.manifest.source_markets.get(f"{dataset}:{market}")
        if market_id is None:
            raise ValueError(f"Configure the original {dataset} market {market}")
        return Context(
            vehicle_model_ids=[self._model(code) for code in models],
            market_ids=[market_id],
            country_codes=[],
            engine_types=[],
            as_of=None,
            scope_notes=[
                "Original dataset grouping; country population and engine scope are unknown."
            ],
        )

    def _source(self, path: Path, kind: str) -> str:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        node_id = f"source:{kind}:{digest}"
        if any(node.id == node_id for node in self.bundle.nodes):
            raise ValueError("This source version is already present in the import")
        self.bundle.nodes.append(
            SourceDocument(
                id=node_id,
                entity_key=node_id,
                label=path.name,
                revision=1,
                type="SourceDocument",
                properties={
                    "source_kind": kind,
                    "origin": self.origin,
                    "location": {"kind": "local_file", "path": str(path)},
                    "publisher": None,
                    "publication_date": None,
                    "retrieved_at": self.created_at,
                    "content_hash": f"sha256:{digest}",
                    "source_role": "primary",
                    "independence_group": None,
                    "access": "provided_internal",
                    "quality_notes": [
                        "Values are source-reported; import does not establish representativeness.",
                        "Formula values, when present, use the source workbook's cached results.",
                    ],
                },
            )
        )
        return node_id

    def _number(self, cell: Cell | MergedCell, source_cell: Cell | MergedCell) -> float | None:
        value = cell.value
        if value is None:
            if source_cell.data_type == "f":
                raise ValueError(f"Formula lacks a cached value at {cell.coordinate}")
            self.report.missing_cells += 1
            return None
        if (
            isinstance(value, bool)
            or not isinstance(value, (int, float))
            or not math.isfinite(value)
        ):
            raise ValueError(f"Expected a finite source number at {cell.coordinate}")
        return float(value)

    def _evidence(
        self,
        source_id: str,
        sheet: str,
        cell: str,
        context: Context,
        measurement: Measurement,
        kind: str,
        notes: list[str],
    ) -> str:
        digest = hashlib.sha256(f"{source_id}\0{sheet}\0{cell}".encode()).hexdigest()
        node_id = f"evidence:{digest}"
        self.bundle.nodes.append(
            Evidence(
                id=node_id,
                entity_key=node_id,
                label=f"{sheet}!{cell}",
                revision=1,
                type="Evidence",
                properties={
                    "evidence_kind": kind,
                    "record_key": f"{sheet}!{cell}",
                    "dedup_key": f"{source_id}:{sheet}!{cell}",
                    "source_channel": None,
                    "locator": {"kind": "spreadsheet", "sheet": sheet, "range": cell},
                    "raw_text": None,
                    "measurement": measurement,
                    "raw_classifications": [],
                    "context": context,
                    "limitations": notes,
                },
            )
        )
        self.bundle.edges.append(
            Edge(
                id=f"edge:source:{digest}",
                type="IN_SOURCE",
                source_id=node_id,
                target_id=source_id,
                rationale="Original workbook value cell; source version is identified by its content hash.",
                context_transfer=None,
            )
        )
        return node_id

    def _measurement(
        self,
        value: float,
        *,
        metric: str,
        unit: str,
        statistic: str,
        year: int,
        question: str | None = None,
        category: str | None = None,
        notes: list[str] | None = None,
    ) -> Measurement:
        return Measurement(
            metric=metric,
            value=value,
            value_range=None,
            unit=unit,
            statistic=statistic,
            temporal_status="source_reported_future"
            if year > self.bundle.as_of.year
            else "historical",
            period={"label": str(year), "start": None, "end": None},
            sample_size=None,
            denominator=None,
            scale=None,
            comparison_basis=None,
            population=None,
            test_standard=None,
            conditions=[],
            question=question,
            response_category=category,
            notes=notes or [],
        )

    def import_sales(self, path: Path) -> None:
        source_id = self._source(path, "sales_dataset")
        values = load_workbook(path, data_only=True, read_only=False)
        formulas = load_workbook(path, data_only=False, read_only=False)
        try:
            groups = {"F70": ["F70"], "G60_G68": ["G60", "G68"], "G70": ["G70"]}
            if set(values.sheetnames) != set(groups):
                raise ValueError("Sales workbook sheets do not match the declared source mapping")
            for sheet_name, models in groups.items():
                sheet = values[sheet_name]
                headers = [cell.value for cell in sheet[1]][:5]
                if headers != [
                    "market",
                    "market_code",
                    "volume_2024",
                    "volume_2025",
                    "volume_2030",
                ]:
                    raise ValueError(
                        "Sales workbook headers do not match the declared source mapping"
                    )
                for row in range(2, sheet.max_row + 1):
                    if all(sheet.cell(row, column).value is None for column in range(1, 6)):
                        continue
                    market = sheet.cell(row, 2).value
                    context = self._context(models, "sales", market)
                    for column, year in ((3, 2024), (4, 2025), (5, 2030)):
                        cell = sheet.cell(row, column)
                        number = self._number(cell, formulas[sheet_name].cell(row, column))
                        if number is None:
                            continue
                        if number < 0 or not number.is_integer():
                            raise ValueError(
                                f"Sales count must be a nonnegative integer at {cell.coordinate}"
                            )
                        notes = ["Forecast method is unknown for source-reported future values."]
                        if len(models) > 1:
                            notes.append(
                                "Combined G60/G68 total; no model-specific allocation is supplied."
                            )
                        evidence_id = self._evidence(
                            source_id,
                            sheet_name,
                            cell.coordinate,
                            context,
                            self._measurement(
                                number,
                                metric="sales_volume",
                                unit="vehicles",
                                statistic="count",
                                year=year,
                            ),
                            "sales_statistic",
                            notes,
                        )
                        if len(models) > 1:
                            self.config.sales_groups[evidence_id] = sheet_name
                        self.report.sales_measurements += 1
        finally:
            values.close()
            formulas.close()

    def _study_value(
        self,
        source_id: str,
        sheet: str,
        cell: Cell,
        number: float,
        model: str,
        market: str,
        question: str,
        question_cell: str,
        category: str | None = None,
        total_cell: str | None = None,
    ) -> None:
        self._questions.add(question)
        notes = [
            f"Original question at {sheet}!{question_cell}.",
            "Sample size, sampling frame, and methodology are unknown.",
        ]
        if total_cell:
            notes.append(
                f"Sample total at {sheet}!{total_cell} is normalized, not a respondent count."
            )
        if category is None:
            notes.append("Rating-scale bounds are unknown; no sentiment conversion is applied.")
        else:
            notes.append(f"Original response category at {sheet}!A{cell.row}.")
        context = self._context([model], "studies", market)
        measurement = self._measurement(
            number,
            metric="survey_response" if category is not None else "survey_mean",
            unit="proportion" if category is not None else "unknown_rating_scale",
            statistic="proportion" if category is not None else "mean",
            year=2025,
            question=question,
            category=category,
            notes=notes,
        )
        evidence_id = self._evidence(
            source_id, sheet, cell.coordinate, context, measurement, "study_statistic", notes
        )
        if category is None:
            self.report.study_means += 1
        else:
            self.report.study_shares += 1
        mapping = self.config.study_mappings.get(question)
        if mapping is None or mapping.status != "reviewed":
            return
        signal_id = evidence_id.replace("evidence:", "signal:study:")
        statement = f"Source-reported {measurement.statistic} for {question}"
        if category is not None:
            statement += f"; response category {category}"
        self.bundle.nodes.append(
            Signal(
                id=signal_id,
                entity_key=signal_id,
                label=cell.coordinate,
                revision=1,
                type="Signal",
                properties={
                    "signal_kind": "study_finding",
                    "statement": statement,
                    "epistemic_status": "publisher_claim",
                    "sentiment": "not_applicable",
                    "context": context,
                    "public_context": None,
                    "capability": None,
                    "extraction_confidence": {"level": "unassessed", "reason": mapping.rationale},
                    "provenance": {
                        "actor_type": "software",
                        "actor_id": "bmwow-workbook-import",
                        "method": "rule_mapping",
                        "created_at": self.created_at,
                        "model": None,
                        "prompt_version": None,
                    },
                },
            )
        )
        for relation, target in (("DERIVED_FROM", evidence_id), ("CONCERNS", mapping.feature_id)):
            self.bundle.edges.append(
                Edge(
                    id=f"edge:{relation.lower()}:{signal_id[len('signal:') :]}",
                    type=relation,
                    source_id=signal_id,
                    target_id=target,
                    rationale=mapping.rationale,
                    context_transfer=None,
                )
            )
        self.report.reviewed_study_findings += 1

    def import_studies(self, path: Path) -> None:
        source_id = self._source(path, "customer_study")
        values = load_workbook(path, data_only=True, read_only=False)
        formulas = load_workbook(path, data_only=False, read_only=False)
        try:
            if set(values.sheetnames) != {"CN_EU_2025", "US_2025"}:
                raise ValueError("Study workbook sheets do not match the declared source mapping")
            mean_sheet = values["CN_EU_2025"]
            for coordinate, labels in (("B2", {"cn", "china"}), ("D2", {"eu", "europe"})):
                if str(mean_sheet[coordinate].value).strip().casefold() not in labels:
                    raise ValueError(
                        "Study region headers do not match the declared column mapping"
                    )
            columns = (
                (2, "G68", "CN"),
                (3, "G70", "CN"),
                (4, "F70", "EU"),
                (5, "G60", "EU"),
                (6, "G70", "EU"),
            )
            for column, model, _ in columns:
                if model not in str(mean_sheet.cell(3, column).value):
                    raise ValueError("Study model headers do not match the declared column mapping")
            question, question_cell = None, None
            for row in range(3, mean_sheet.max_row + 1):
                label = mean_sheet.cell(row, 1).value
                if not isinstance(label, str) or not label.strip():
                    continue
                if label.strip() != "Mean":
                    question, question_cell = label, f"A{row}"
                    continue
                if question is None:
                    raise ValueError("Study mean has no preceding original question")
                for column, model, market in columns:
                    cell = mean_sheet.cell(row, column)
                    number = self._number(cell, formulas[mean_sheet.title].cell(row, column))
                    if number is not None:
                        self._study_value(
                            source_id,
                            mean_sheet.title,
                            cell,
                            number,
                            model,
                            market,
                            question,
                            question_cell,
                        )
            sheet = values["US_2025"]
            for column, model in ((3, "G60"), (4, "G70")):
                if model not in str(sheet.cell(1, column).value):
                    raise ValueError(
                        "US study model headers do not match the declared column mapping"
                    )
            question, question_cell, total_cell = None, None, None
            for row in range(1, sheet.max_row + 1):
                label = sheet.cell(row, 1).value
                marker = sheet.cell(row, 2).value
                if label is None and marker is None:
                    continue
                if marker is None and isinstance(label, str):
                    question, question_cell, total_cell = label, f"A{row}", None
                    continue
                if not isinstance(marker, str) or marker.strip() != "% of column, sample base":
                    raise ValueError(f"Unrecognized study row layout at US_2025!A{row}")
                if question is None or not isinstance(label, str) or not label.strip():
                    raise ValueError("Study share lacks an original question or response category")
                for column, model in ((3, "G60"), (4, "G70")):
                    cell = sheet.cell(row, column)
                    number = self._number(cell, formulas[sheet.title].cell(row, column))
                    if number is None:
                        continue
                    if label.strip() == "Sample total":
                        if number != 1:
                            raise ValueError(
                                f"Unexpected normalized sample total at {cell.coordinate}"
                            )
                        self.report.normalized_totals += 1
                        continue
                    if not 0 <= number <= 1:
                        raise ValueError(f"Study proportion outside [0,1] at {cell.coordinate}")
                    self._study_value(
                        source_id,
                        sheet.title,
                        cell,
                        number,
                        model,
                        "US",
                        question,
                        question_cell,
                        category=label,
                        total_cell=f"{get_column_letter(column)}{total_cell}"
                        if total_cell
                        else None,
                    )
                if label.strip() == "Sample total":
                    total_cell = row
        finally:
            values.close()
            formulas.close()

    def finish(self) -> tuple[GraphBundle, KnowledgeConfig, ImportReport]:
        self.report.questions = sorted(self._questions)
        self.report.unmapped_questions = sum(
            question not in self.config.study_mappings
            or self.config.study_mappings[question].status != "reviewed"
            for question in self._questions
        )
        errors = validate_bundle(self.bundle)
        if errors:
            raise ValueError(f"Imported graph failed validation ({len(errors)} errors)")
        KnowledgeService(InMemoryRepository(self.bundle), self.config)
        return (
            self.bundle.model_copy(deep=True),
            self.config.model_copy(deep=True),
            self.report.model_copy(deep=True),
        )
